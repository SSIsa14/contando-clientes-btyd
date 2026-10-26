"""Reconciliación independiente desde Excel y tablas guardadas, sin reajustar.

No usa limpiar ni historias_retail para su oráculo de conteos/netos. Asistencia de
Codex (OpenAI), revisión humana pendiente. Uso: python scripts/validar_retail.py
"""
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, localcontext
import hashlib
import json
from pathlib import Path
import sys

import nbformat
import numpy as np
import pandas as pd
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
from bgnbd import esperado_Y, loglik_bgnbd, NOMBRES


def sha(ruta):
    return hashlib.sha256(Path(ruta).read_bytes()).hexdigest()


def leer(nombre):
    return json.loads((RAIZ / "resultados" / nombre).read_text(encoding="utf-8"))


def validar():
    ruta = RAIZ / "data/Online Retail.xlsx"
    huella = "43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d"
    assert sha(ruta) == huella
    audit = leer("retail_limpieza_auditoria.json")
    pro = leer("ajuste_retail_proveniencia.json")
    hold = leer("retail_holdout.json")
    params = leer("params_retail.json")
    assert set(params) == set(NOMBRES)
    assert all(type(v) is float and np.isfinite(v) and v > 0 for v in params.values())
    assert params == pro["parametros"] and pro["sustitucion_limpiador"]["original_recuperada"] is False
    assert sha(RAIZ / "resultados/params_cdnow.json") == pro["params_cdnow_sha256_sin_modificar"]
    for nombre, esperado in {**pro["codigo_sha256"], **pro["artefactos_sha256"]}.items():
        assert sha(RAIZ / nombre) == esperado, nombre

    # Oráculo separado: openpyxl y sumas Decimal, no la escala entera del limpiador.
    libro = load_workbook(ruta, read_only=True, data_only=True)
    filas = libro["Online Retail"].iter_rows(values_only=True)
    cabecera = next(filas)
    assert tuple(cabecera) == ("InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID", "Country")
    netos, n_lineas = defaultdict(Decimal), defaultdict(int)
    facturas, facturas_id = set(), set()
    n = sin_id = devueltas = canceladas = 0
    with localcontext() as contexto:
        contexto.prec = 50
        for invoice, stock, desc, cantidad, fecha, precio, cliente, pais in filas:
            n += 1
            facturas.add(str(invoice))
            assert isinstance(fecha, datetime)
            q, p = Decimal(str(cantidad)), Decimal(str(precio))
            assert q.is_finite() and p.is_finite()
            if cliente is None:
                sin_id += 1
                continue
            clave = (str(int(cliente)), fecha.date())
            netos[clave] += q * p
            n_lineas[clave] += 1
            facturas_id.add(str(invoice))
            devueltas += q < 0
            canceladas += str(invoice).upper().startswith("C")
    libro.close()
    assert n == audit["filas_originales"] == 541909
    assert sin_id == audit["filas_sin_customerid"]
    assert len(facturas) == audit["facturas_unicas_originales"]
    assert len(facturas_id) == audit["facturas_unicas_con_customerid"]
    assert devueltas == audit["filas_quantity_negativa_conservadas"]
    assert canceladas == audit["filas_cancelacion_conservadas"]
    dias = pd.read_csv(RAIZ / "resultados/retail_dias.csv", dtype=str)
    assert len(netos) == len(dias) == audit["dias_cliente_totales"]
    assert not dias.duplicated(["cliente", "fecha"]).any()
    for fila in dias.itertuples():
        clave = (fila.cliente, pd.Timestamp(fila.fecha).date())
        assert Decimal(fila.venta_neta_exacta) == netos[clave]
        assert int(fila.lineas) == n_lineas[clave]
        assert (fila.califica == "True") == (netos[clave] > 0)
    assert sum(v == 0 for v in netos.values()) == audit["dias_neto_cero"]
    assert sum(v < 0 for v in netos.values()) == audit["dias_neto_negativo"]
    positivos = [k for k, v in netos.items() if v > 0]
    assert len(positivos) == audit["transacciones_calificantes_todo_archivo"]

    corte, final = datetime(2011, 6, 30).date(), datetime(2011, 12, 8).date()
    calibracion, futuro = defaultdict(list), defaultdict(list)
    for cliente, fecha in positivos:
        if fecha <= corte:
            calibracion[cliente].append(fecha)
        elif fecha <= final:
            futuro[cliente].append(fecha)
    clientes = pd.read_csv(RAIZ / "resultados/retail_clientes.csv", dtype={"cliente": str},
                           converters={k: float for k in ["tx", "T", "esperado_23"]})
    assert clientes.cliente.is_unique and set(clientes.cliente) == set(calibracion)
    for fila in clientes.itertuples():
        fechas = sorted(calibracion[fila.cliente])
        assert fila.x == len(fechas) - 1
        assert fila.x_star == len(futuro[fila.cliente])
        assert pd.Timestamp(fila.adquisicion).date() == fechas[0]
        np.testing.assert_allclose(fila.tx, (fechas[-1] - fechas[0]).days / 7, atol=1e-14)
        np.testing.assert_allclose(fila.T, (corte - fechas[0]).days / 7, atol=1e-14)
    assert len(clientes) == hold["cohorte"]["clientes"]
    assert (final - corte).days / 7 == hold["tau_semanas"] == 23
    real_oraculo = sum(len(futuro[c]) for c in calibracion)
    assert real_oraculo == clientes.x_star.sum() == hold["real_total"]
    x, tx, T = (clientes[k].to_numpy(float) for k in ["x", "tx", "T"])
    pred = esperado_Y(23, x, tx, T, params)
    np.testing.assert_allclose(clientes.esperado_23, pred, atol=1e-12, rtol=1e-12)
    np.testing.assert_allclose(loglik_bgnbd(np.log([params[k] for k in NOMBRES]), x, tx, T), pro["ll_sumada"], atol=1e-8, rtol=0)
    total = float(clientes.esperado_23.sum())
    np.testing.assert_allclose(total, hold["pronostico_total"], atol=1e-8, rtol=0)
    np.testing.assert_allclose(total - real_oraculo, hold["error_firmado"], atol=1e-8, rtol=0)
    np.testing.assert_allclose(abs(total - real_oraculo), hold["error_absoluto"], atol=1e-8, rtol=0)
    np.testing.assert_allclose(100 * (total - real_oraculo) / real_oraculo, hold["error_porcentual"], atol=1e-10, rtol=0)

    diario = pd.read_csv(RAIZ / "resultados/retail_diario.csv", parse_dates=["fecha"])
    assert len(diario) == 161 and diario.fecha.is_unique
    por_fecha = defaultdict(int)
    for cliente in calibracion:
        for fecha in futuro[cliente]:
            por_fecha[fecha] += 1
    assert diario.real.tolist() == [por_fecha[f.date()] for f in diario.fecha]
    np.testing.assert_array_equal(diario.real.cumsum(), diario.real_acumulado)
    np.testing.assert_allclose(np.diff(np.r_[0, diario.pronostico_acumulado]), diario.pronostico, atol=1e-10)
    assert diario.fecha.min() == pd.Timestamp("2011-07-01")
    assert diario.fecha.max() == pd.Timestamp("2011-12-08")
    for nombre, col in [("retail_semanal.csv", "semana_holdout"), ("retail_mensual.csv", "mes")]:
        tabla = pd.read_csv(RAIZ / "resultados" / nombre)
        agrupado = diario.groupby(col).agg(real=("real", "sum"), pronostico=("pronostico", "sum"), dias=("fecha", "size"))
        np.testing.assert_array_equal(tabla.real, agrupado.real)
        np.testing.assert_allclose(tabla.pronostico, agrupado.pronostico, atol=1e-9)
        np.testing.assert_allclose(tabla.real_por_dia, tabla.real / tabla.dias, atol=1e-12)
        assert tabla.real.sum() == real_oraculo
        np.testing.assert_allclose(tabla.pronostico.sum(), total, atol=1e-8)
        if nombre == "retail_semanal.csv":
            assert len(tabla) == 23 and tabla.dias.eq(7).all()
        else:
            assert tabla.dias.tolist() == [31, 31, 30, 31, 30, 8]

    nb = nbformat.read(RAIZ / "notebooks/05_online_retail.ipynb", as_version=4)
    nbformat.validate(nb)
    codigo = [c for c in nb.cells if c.cell_type == "code"]
    assert [c.execution_count for c in codigo] == list(range(1, len(codigo) + 1))
    assert not any(o.output_type == "error" for c in codigo for o in c.outputs)
    ajuste = next(c for c in codigo if c.id == "p5-ajuste")
    texto = "".join(o.get("text", "") for o in ajuste.outputs)
    linea = next(l for l in texto.splitlines() if l.startswith("Parámetros completos: "))
    assert json.loads(linea.split(": ", 1)[1]) == params
    assert f"LL sumada: {pro['ll_sumada']!r}" in texto
    for f in ["parte5_holdout_acumulado.png", "parte5_holdout_semanal.png", "parte5_meses_por_dia.png", "parte5_historias.png"]:
        assert (RAIZ / "figs" / f).stat().st_size > 10000
    assert sha(ruta) == huella
    informe = {
        "estado": "auditoría independiente aprobada", "raw_sha256": huella,
        "notebook_sha256": sha(RAIZ / "notebooks/05_online_retail.ipynb"),
        "celdas_codigo_ejecutadas": len(codigo), "oraculo": "openpyxl + Decimal; sin limpiar/resumen_clientes",
        "dias_netos_reconciliados": len(netos), "clientes_reconciliados": len(clientes),
        "real_holdout": real_oraculo, "pronostico": total,
        "verificaciones": ["netos exactos y multiplicidad de líneas", "conteos de líneas/facturas",
            "cohorte y adquisición calificante", "ventanas y 23 semanas", "ceros de holdout",
            "sumas cliente/diario/semanal/mensual", "errores firmados/absolutos/porcentuales",
            "fuentes y artefactos por SHA-256", "notebook sin errores guardados"],
        "alcance": "Parte 5; revisión humana, presentación e integración final pendientes",
    }
    (RAIZ / "resultados/validacion_retail.json").write_text(
        json.dumps(informe, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps(informe, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    validar()

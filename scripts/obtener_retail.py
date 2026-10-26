"""Recupera/verifica el Excel original UCI 352 e inventaría su esquema, sin limpiar.

Asistencia de implementación: Codex (OpenAI); revisión humana pendiente.
Uso desde la raíz: python scripts/obtener_retail.py [--zip archivo.zip]
Un archivo existente con otra huella nunca se sobrescribe.
"""
import argparse
from collections import Counter
from datetime import datetime
import hashlib
import io
import json
from pathlib import Path
from urllib.request import urlopen
import zipfile

from openpyxl import load_workbook

PAGINA = "https://archive.ics.uci.edu/dataset/352/online+retail"
DESCARGA = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
SHA_XLSX = "43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d"
SHA_ZIP = "f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b"


def recuperar(archivo_zip=None):
    raiz = Path(__file__).resolve().parents[1]
    destino = raiz / "data/Online Retail.xlsx"
    if archivo_zip is not None or not destino.exists():
        if archivo_zip is None:
            with urlopen(DESCARGA, timeout=120) as respuesta:
                paquete = respuesta.read()
        else:
            paquete = Path(archivo_zip).read_bytes()
        if hashlib.sha256(paquete).hexdigest() != SHA_ZIP:
            raise ValueError("Cambió la huella del ZIP oficial; investigar antes de usarlo")
        with zipfile.ZipFile(io.BytesIO(paquete)) as z:
            original = z.read("Online Retail.xlsx")
        if hashlib.sha256(original).hexdigest() != SHA_XLSX:
            raise ValueError("Huella inesperada del Excel original")
        if not destino.exists():
            destino.write_bytes(original)
    if hashlib.sha256(destino.read_bytes()).hexdigest() != SHA_XLSX:
        raise ValueError("El Excel existente difiere del original; no se sobrescribe")

    libro = load_workbook(destino, read_only=True, data_only=True)
    if libro.sheetnames != ["Online Retail"]:
        raise ValueError(f"Hojas inesperadas: {libro.sheetnames}")
    hoja = libro["Online Retail"]
    filas = hoja.iter_rows(values_only=True)
    columnas = list(next(filas))
    esperado = ["InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate",
                "UnitPrice", "CustomerID", "Country"]
    if columnas != esperado:
        raise ValueError(f"Esquema inesperado: {columnas}")
    n = 0
    nulos = Counter({k: 0 for k in columnas})
    tipos = {k: Counter() for k in columnas}
    minimo, maximo = None, None
    negativas = cancelaciones = no_positivo_precio = 0
    for fila in filas:
        n += 1
        for clave, valor in zip(columnas, fila):
            nulos[clave] += valor is None
            tipos[clave][type(valor).__name__] += 1
        fecha = fila[4]
        if not isinstance(fecha, datetime):
            raise ValueError("InvoiceDate contiene un valor sin fecha")
        minimo = fecha if minimo is None else min(minimo, fecha)
        maximo = fecha if maximo is None else max(maximo, fecha)
        negativas += fila[3] < 0
        cancelaciones += str(fila[0]).upper().startswith("C")
        no_positivo_precio += fila[5] <= 0
    libro.close()
    if n != 541909:
        raise ValueError(f"Número inesperado de filas: {n}")
    informe = {
        "conjunto": "Online Retail (UCI 352), no Online Retail II",
        "autor_datos": "Daqing Chen", "doi": "10.24432/C5BW33",
        "pagina_oficial": PAGINA, "enlace_download_oficial": DESCARGA,
        "licencia": "CC BY 4.0", "licencia_url": "https://creativecommons.org/licenses/by/4.0/",
        "archivo": "data/Online Retail.xlsx", "bytes": destino.stat().st_size,
        "sha256": SHA_XLSX, "zip_sha256_referencia": SHA_ZIP,
        "hoja": "Online Retail", "filas": n, "columnas": columnas,
        "tipos_celdas_excel": {k: dict(v) for k, v in tipos.items()},
        "nulos": dict(nulos), "fecha_min": minimo.isoformat(), "fecha_max": maximo.isoformat(),
        "filas_quantity_negativa": negativas, "filas_invoice_c": cancelaciones,
        "filas_unitprice_no_positivo": no_positivo_precio,
        "transformaciones": "ninguna; original conservado sin limpieza ni agregación",
        "reproducir": "python scripts/obtener_retail.py",
        "limpiar_original": "no recuperada; reutilización pendiente",
    }
    (raiz / "data/online_retail_origen.json").write_text(
        json.dumps(informe, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps(informe, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--zip", help="ZIP oficial ya descargado; también se verifica")
    recuperar(parser.parse_args().zip)
